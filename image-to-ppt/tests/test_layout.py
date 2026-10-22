import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from image_to_ppt import exporters, imaging
from image_to_ppt.layout import (
    MODE_FIT,
    MODE_HEIGHT_FIRST,
    MODE_WIDTH_FIRST,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    compute_layout,
)

A4_PORTRAIT = (2480, 3508)
A4_LANDSCAPE = (3508, 2480)
WIDE = (4000, 1000)


# --- 排版計算 ---------------------------------------------------------------


def test_height_first_matches_spec():
    rect = compute_layout(*A4_PORTRAIT, mode=MODE_HEIGHT_FIRST)
    assert rect.h == SLIDE_HEIGHT
    assert rect.y == 0
    assert rect.w == pytest.approx(A4_PORTRAIT[0] * SLIDE_HEIGHT / A4_PORTRAIT[1])
    assert rect.x == pytest.approx((SLIDE_WIDTH - rect.w) / 2)


def test_height_first_clips_wide_image():
    rect = compute_layout(*WIDE, mode=MODE_HEIGHT_FIRST)
    assert rect.h == SLIDE_HEIGHT
    assert rect.w > SLIDE_WIDTH
    assert rect.x < 0


def test_fit_keeps_inside_and_centers():
    for size in (A4_PORTRAIT, A4_LANDSCAPE, WIDE):
        rect = compute_layout(*size, mode=MODE_FIT)
        assert rect.w <= SLIDE_WIDTH + 1e-6
        assert rect.h <= SLIDE_HEIGHT + 1e-6
        assert rect.x == pytest.approx((SLIDE_WIDTH - rect.w) / 2)
        assert rect.y == pytest.approx((SLIDE_HEIGHT - rect.h) / 2)


def test_fit_preserves_aspect_ratio():
    rect = compute_layout(*A4_PORTRAIT, mode=MODE_FIT)
    assert rect.w / rect.h == pytest.approx(A4_PORTRAIT[0] / A4_PORTRAIT[1])


def test_width_first_matches_spec():
    rect = compute_layout(*A4_PORTRAIT, mode=MODE_WIDTH_FIRST)
    assert rect.w == SLIDE_WIDTH
    assert rect.x == 0
    assert rect.h == pytest.approx(A4_PORTRAIT[1] * SLIDE_WIDTH / A4_PORTRAIT[0])
    assert rect.y == pytest.approx((SLIDE_HEIGHT - rect.h) / 2)


def test_width_first_clips_portrait():
    rect = compute_layout(*A4_PORTRAIT, mode=MODE_WIDTH_FIRST)
    assert rect.h > SLIDE_HEIGHT
    assert rect.y < 0


def test_all_modes_preserve_aspect_ratio():
    for mode in (MODE_HEIGHT_FIRST, MODE_FIT, MODE_WIDTH_FIRST):
        rect = compute_layout(*A4_LANDSCAPE, mode=mode)
        assert rect.w / rect.h == pytest.approx(A4_LANDSCAPE[0] / A4_LANDSCAPE[1])


def test_invalid_mode_raises():
    with pytest.raises(ValueError):
        compute_layout(100, 100, mode="nope")


def test_invalid_size_raises():
    with pytest.raises(ValueError):
        compute_layout(0, 100)


# --- 測試輔助 ---------------------------------------------------------------


def _png_bytes(width: int, height: int) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "red").save(buffer, format="PNG")
    return buffer.getvalue()


def _make_pdf(path, pages: int) -> None:
    import pymupdf

    document = pymupdf.open()
    for _ in range(pages):
        document.new_page()
    document.save(str(path))
    document.close()


# --- 輸入 -------------------------------------------------------------------


def test_load_image_reads_size_and_bytes(tmp_path):
    path = tmp_path / "pic.png"
    path.write_bytes(_png_bytes(640, 480))

    image = imaging.load_image(path)
    assert (image.width, image.height) == (640, 480)
    assert image.name == "pic.png"
    assert image.data


def test_collect_slide_images_expands_pdf_pages(tmp_path):
    image_path = tmp_path / "pic.png"
    image_path.write_bytes(_png_bytes(640, 480))
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 3)

    images = imaging.collect_slide_images([image_path, pdf_path])
    assert images[0].name == "pic.png"
    assert len(images) == 4
    assert all(image.data for image in images)


def test_collect_slide_images_rejects_unknown(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hi", encoding="utf-8")
    with pytest.raises(ValueError):
        imaging.collect_slide_images([path])


def test_preview_only_renders_first_page(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 5)

    image = imaging.preview_slide_image(pdf_path)
    assert image is not None
    assert image.name == "doc.pdf - 第 1 頁"


def test_preview_missing_file_returns_none(tmp_path):
    assert imaging.preview_slide_image(tmp_path / "missing.pdf") is None


# --- 輸出 -------------------------------------------------------------------


def test_build_presentation_slide_dimensions_and_count(tmp_path):
    images = [
        imaging.SlideImage(_png_bytes(800, 1200), 800, 1200, "a"),
        imaging.SlideImage(_png_bytes(1600, 900), 1600, 900, "b"),
    ]
    output = tmp_path / "out.pptx"
    seen: list[tuple[int, int]] = []

    exporters.build_presentation(
        images, output, mode=MODE_HEIGHT_FIRST, progress=lambda d, t: seen.append((d, t))
    )

    assert output.exists()
    prs = Presentation(str(output))
    assert len(prs.slides) == 2
    assert prs.slide_width == Emu(exporters.SLIDE_WIDTH_EMU)
    assert prs.slide_height == Emu(exporters.SLIDE_HEIGHT_EMU)
    assert seen == [(1, 2), (2, 2)]

    picture = prs.slides[0].shapes[0]
    expected = compute_layout(800, 1200, MODE_HEIGHT_FIRST)
    assert picture.height == Emu(round(expected.h * exporters.PX_TO_EMU))


def test_export_pdf_images_names_by_page_number(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 12)
    out_dir = tmp_path / "images"

    seen: list[tuple[int, int]] = []
    files = exporters.export_pdf_images(
        pdf_path, out_dir, dpi=72, image_format="PNG", progress=lambda d, t: seen.append((d, t))
    )

    assert [f.name for f in files] == [f"{i:03d}.png" for i in range(1, 13)]
    assert all(f.exists() for f in files)
    assert seen == [(i, 12) for i in range(1, 13)]


def test_export_pdf_images_jpeg_format(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 2)
    files = exporters.export_pdf_images(
        pdf_path, tmp_path / "out", dpi=72, image_format="JPEG"
    )
    assert [f.suffix for f in files] == [".jpg", ".jpg"]


def test_export_pdf_images_rejects_empty_pdf(tmp_path, monkeypatch):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 1)
    monkeypatch.setattr(exporters, "iter_pdf_pages", lambda *a, **k: iter(()))
    with pytest.raises(ValueError):
        exporters.export_pdf_images(pdf_path, tmp_path / "out", dpi=72)


def test_iter_pdf_pages_reports_number_and_total(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 4)

    pages = list(imaging.iter_pdf_pages(pdf_path, dpi=72))
    assert [page.number for page in pages] == [1, 2, 3, 4]
    assert all(page.total == 4 for page in pages)


def test_expand_refs_splits_pdf_pages(tmp_path):
    image_path = tmp_path / "pic.png"
    image_path.write_bytes(_png_bytes(640, 480))
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 3)

    refs = imaging.expand_refs([image_path, pdf_path])
    assert refs[0] == imaging.SlideRef(image_path)
    assert [ref.page for ref in refs[1:]] == [1, 2, 3]
    assert all(ref.path == pdf_path for ref in refs[1:])


def test_slide_image_renders_selected_page(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 5)

    image = imaging.slide_image(imaging.SlideRef(pdf_path, 3), dpi=72)
    assert image.name == "doc.pdf - 第 3 頁"


def test_preview_ref_specific_page_and_missing(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    _make_pdf(pdf_path, 5)

    image = imaging.preview_ref(imaging.SlideRef(pdf_path, 2))
    assert image is not None and image.name == "doc.pdf - 第 2 頁"
    assert imaging.preview_ref(imaging.SlideRef(tmp_path / "missing.pdf", 1)) is None


def test_slide_ref_label(tmp_path):
    assert imaging.SlideRef(tmp_path / "a.png").label == "a.png"
    assert imaging.SlideRef(tmp_path / "a.pdf", 2).label == "a.pdf 第 2 頁"


def test_slide_is_exactly_16_9():
    assert exporters.SLIDE_WIDTH_EMU / exporters.SLIDE_HEIGHT_EMU == pytest.approx(16 / 9)
