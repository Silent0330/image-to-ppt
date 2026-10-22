"""輸出：PowerPoint 簡報，以及 PDF 轉圖片資料夾。"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Callable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from .imaging import SlideImage, iter_pdf_pages
from .layout import MODE_HEIGHT_FIRST, SLIDE_WIDTH, compute_layout

EMU_PER_INCH = 914400
SLIDE_HEIGHT_INCHES = 7.5
SLIDE_HEIGHT_EMU = round(SLIDE_HEIGHT_INCHES * EMU_PER_INCH)
SLIDE_WIDTH_EMU = round(SLIDE_HEIGHT_EMU * 16 / 9)
SLIDE_WIDTH_INCHES = SLIDE_WIDTH_EMU / EMU_PER_INCH
PX_TO_EMU = SLIDE_WIDTH_EMU / SLIDE_WIDTH

OUTPUT_IMAGE_FORMATS = {"PNG": "png", "JPEG": "jpg"}
DEFAULT_BACKGROUND = "FFFFFF"

Progress = Callable[[int, int], None]


def _emu(pixels: float) -> Emu:
    return Emu(round(pixels * PX_TO_EMU))


def _add_slide(prs: Presentation, background: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(background.upper())
    return slide


def build_presentation(
    images: Sequence[SlideImage],
    output_path: str | Path,
    *,
    mode: str = MODE_HEIGHT_FIRST,
    background: str = DEFAULT_BACKGROUND,
    progress: Progress | None = None,
) -> Path:
    """每張圖片對應一張投影片，輸出成 .pptx。"""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(SLIDE_HEIGHT_EMU)

    total = len(images)
    for index, image in enumerate(images, start=1):
        slide = _add_slide(prs, background)
        rect = compute_layout(image.width, image.height, mode)
        picture = slide.shapes.add_picture(
            io.BytesIO(image.data),
            _emu(rect.x),
            _emu(rect.y),
            width=_emu(rect.w),
            height=_emu(rect.h),
        )
        picture.name = image.name
        if progress is not None:
            progress(index, total)

    output_path = Path(output_path)
    prs.save(str(output_path))
    return output_path


def export_pdf_images(
    pdf_path: str | Path,
    output_dir: str | Path,
    *,
    dpi: int,
    image_format: str = "PNG",
    progress: Progress | None = None,
) -> list[Path]:
    """將 PDF 每一頁輸出成圖片，檔名為零補位的頁碼（例如 001.png）。"""
    pdf_path = Path(pdf_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    extension = OUTPUT_IMAGE_FORMATS[image_format.upper()]

    digits = 3
    total = 0
    written: list[Path] = []
    for page in iter_pdf_pages(pdf_path, dpi):
        if page.number == 1:
            total = page.total
            digits = max(3, len(str(total)))
        target = output_dir / f"{page.number:0{digits}d}.{extension}"
        page.pixmap.save(str(target))
        written.append(target)
        if progress is not None:
            progress(page.number, total)

    if not written:
        raise ValueError("PDF 沒有可輸出的頁面。")
    return written
